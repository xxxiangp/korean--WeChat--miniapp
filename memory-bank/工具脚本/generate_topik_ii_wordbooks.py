#!/usr/bin/env python3
"""Generate auditable TOPIK II intermediate and advanced wordbooks."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import tempfile
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = ROOT / "memory-bank" / "数据文件"
DOC_DIR = ROOT / "memory-bank" / "产品文档"
TODAY = date.today().isoformat()

KRDICT_PATH = DATA_DIR / "topik_ii_krdict_metadata.json"
AUDIT_JSON_PATH = DATA_DIR / "topik_ii_metadata_audit.json"
AUDIT_MD_PATH = DOC_DIR / "TOPIK中高级词汇元数据审计.md"

LEVEL_CONFIGS = {
    "topik_intermediate": {
        "level_cn": "TOPIK中级",
        "source_pool": DATA_DIR / "topik_intermediate_source_pool.json",
        "expected_total": 3715,
        "description": (
            "TOPIK II 3~4级备考词汇，基于国立国语院学习等级与公开TOPIK II"
            "词表交叉审计；非官方固定考纲"
        ),
        "quiz_description": "TOPIK II 3~4级备考词汇四选一测验数据",
        "docx": DATA_DIR / "TOPIK中级词汇_最终交付版.docx",
    },
    "topik_advanced": {
        "level_cn": "TOPIK高级",
        "source_pool": DATA_DIR / "topik_advanced_source_pool.json",
        "expected_total": 1008,
        "description": (
            "TOPIK II 5~6级备考词汇，基于国立国语院学习等级与公开TOPIK II"
            "词表交叉审计；非官方固定考纲"
        ),
        "quiz_description": "TOPIK II 5~6级备考词汇四选一测验数据",
        "docx": DATA_DIR / "TOPIK高级词汇_最终交付版.docx",
    },
}

POS_CN_MAP = {
    "명사": "名词",
    "동사": "动词",
    "형용사": "形容词",
    "부사": "副词",
    "감탄사": "感叹词",
    "조사": "助词",
    "관형사": "冠形词",
    "대명사": "代词",
    "수사": "数词",
}
POS_PRIORITY = {
    "명사": 0,
    "동사": 1,
    "형용사": 2,
    "부사": 3,
    "관형사": 4,
    "대명사": 5,
    "수사": 6,
    "감탄사": 7,
    "조사": 8,
}
LEGACY_POS_MAP = {
    "명": "명사",
    "동": "동사",
    "형": "형용사",
    "부": "부사",
    "감": "감탄사",
    "조": "조사",
    "관": "관형사",
    "대": "대명사",
    "수": "수사",
    "의": "명사",
    "고": "명사",
}

# The official dictionary marks these contractions as "shortened forms" rather
# than ordinary parts of speech. The values below restore their base-word POS.
POS_OVERRIDES = {
    "걔": "대명사",
    "그래도": "부사",
    "그래야": "부사",
    "그렇게": "부사",
    "어때": "형용사",
    "어떡하다": "동사",
    "어째서": "부사",
    "이래서": "부사",
    "이렇게": "부사",
    "얘": "대명사",
    "쟤": "대명사",
    "저렇게": "부사",
    "싶어지다": "동사",
}

# These 48 candidates are absent as exact headwords from the official Korean
# Basic Dictionary snapshot. 애기 is present but has no Chinese translation.
# Each fallback was reviewed against its source gloss and made concise here.
MEANING_OVERRIDES = {
    "가능해지다": "变得可能",
    "가정교사": "家庭教师",
    "고속도로": "高速公路",
    "국회의원": "国会议员",
    "그려지다": "被画出；被描绘；浮现在脑海中",
    "그제서야": "直到那时才",
    "김포공항": "金浦机场",
    "느껴지다": "被感受到；感觉到",
    "다양해지다": "变得多样",
    "대학로": "大学路（首尔文化艺术街区）",
    "동대문시장": "东大门市场",
    "만들어지다": "被制作；形成",
    "몇십": "几十",
    "문제되다": "成为问题",
    "미국인": "美国人",
    "분명해지다": "变得明确；变得清楚",
    "사이버": "网络；虚拟空间",
    "생활수준": "生活水平",
    "생활환경": "生活环境",
    "세워지다": "被竖立；被建立",
    "세종대왕": "世宗大王",
    "심각해지다": "变得严重",
    "심해지다": "变得严重；加剧",
    "싶어지다": "开始想要",
    "안동": "安东（韩国城市）",
    "약해지다": "变弱",
    "어두워지다": "变暗",
    "어려워지다": "变得困难",
    "여겨지다": "被认为；被视为",
    "여직원": "女职员",
    "윗글": "上文",
    "유능": "才能；有能力",
    "이뤄지다": "实现；达成",
    "인천공항": "仁川国际机场",
    "일본인": "日本人",
    "잊혀지다": "被遗忘",
    "자연현상": "自然现象",
    "전해지다": "被传达；流传",
    "정해지다": "被决定；确定",
    "좋아지다": "变好；改善",
    "중국인": "中国人",
    "지워지다": "被擦除；被抹去",
    "친해지다": "变得亲近",
    "켜지다": "被打开；亮起",
    "행해지다": "被实行；被执行",
    "환경오염": "环境污染",
    "활발해지다": "变得活跃",
    "힘들어하다": "感到困难；感到辛苦",
    "애기": "婴儿；小孩",
    "걔": "那孩子；那家伙",
    "그래도": "即使那样；尽管如此",
    "데": "地方；情况；时候",
    "어떡하다": "怎么办；如何处理",
    "이래서": "因此；这样一来",
    "쟤": "那孩子；那家伙",
    "활짝": "完全敞开地；灿烂地",
    "나름": "各自的方式；各有各的情况",
    "님": "您；先生或女士（敬称）",
    "동양인": "东亚人；东方人",
    "머리칼": "头发",
    "그렇게": "那样；如此",
    "만큼": "像……一样；达到……程度",
    "어때": "怎么样",
    "이렇게": "这样；如此",
    "저렇게": "那样；如此",
    "년생": "……年出生的人",
    "어째서": "为什么；为何",
}

# The first group merges senses confirmed through the official Standard Korean
# Dictionary. The second group condenses long Korean Basic Dictionary glosses
# without changing their substance, so meanings remain readable in the app.
FORCE_MEANING_OVERRIDES = {
    "갑": "甲；甲级；盒；匣；岬角",
    "인": "瘾；人；磷",
    "동화": "同化；童话；铜币",
    "침": "唾液；针；针灸",
    "부처": "佛祖；佛像；部门；夫妻",
    "소요": "所需；骚乱；逍遥；散步",
    "운동량": "运动量；动量",
    "리": "理由；可能性；里；厘",
    "코스모스": "大波斯菊；宇宙",
    "얘": "这孩子；这家伙",
    "곱다": "美丽；细腻；温和",
    "끊어지다": "被切断；中断；断绝；停止",
    "넘어가다": "越过；翻过；转入；超过；被骗；去世",
    "대다": "接触；靠；提供；对准；比较；说出",
    "물리다": "厌烦；被咬；退还；使退下；继承；含住；缴纳",
    "뵈다": "看见；拜见；看起来；给……看",
    "잡히다": "被抓住；被逮捕；被掌握；被确定；结冰；起泡",
    "통하다": "相通；畅通；沟通；通过；行得通",
    "풀리다": "被解开；解除；消除；放松；融化",
    "환하다": "明亮；开阔；清楚；开朗",
    "거칠다": "粗糙；荒凉；粗鲁；猛烈；急促",
    "내놓다": "拿出来；发表；提出；让出；显露",
    "박히다": "被钉住；被嵌入；被印刻；铭记；固定",
    "터지다": "破裂；爆炸；爆发；绽开；突然涌出",
    "깨끗해지다": "变干净；变晴朗；痊愈；变清白",
    "떨어뜨리다": "使掉落；使落榜；降低；耗尽；疏远",
    "트이다": "被打开；畅通；开朗；开窍；放声",
    "높이다": "提高；增高；增强；抬高；尊敬",
    "험하다": "险峻；危险；粗暴；艰难；触目惊心",
    "끊기다": "被切断；中断；断绝；挂断；停运",
    "맡기다": "交付；委托；托管；让……担任；抵押",
    "풀어지다": "被解开；放松；消散；解除；回暖",
    "당하다": "遭受；面临；担当；比得过；相当于",
    "단단하다": "坚硬；结实；牢固；坚定",
    "비치다": "照射；映出；显现；流露；露面；视为",
    "넘어오다": "倒向这边；转移过来；转入；中计；越过来",
    "묶이다": "被捆绑；被限制；被组成；被汇总",
    "무너지다": "倒塌；垮台；崩溃；落空；败北",
    "죽이다": "杀死；熄灭；消除；压低；消磨；使变钝",
    "흐려지다": "变模糊；变浑浊；变阴沉；变暗淡",
    "닿다": "接触；到达；传来；涉及；理解；联系",
    "옮기다": "搬动；迁移；转换；翻译；传播；传染",
    "썩다": "腐烂；腐朽；堕落；被埋没；操心；浪费",
    "늘어지다": "变长；松弛；无精打采；拖延；尽情",
    "가리다": "遮挡；分辨；认生；挑食；自理",
    "떼다": "取下；扣除；断绝；转移；批发；戒除",
    "벗기다": "脱下；卸下；剥去；揭开；洗刷；搜刮",
    "다루다": "处理；经营；操作；对待；管教；探讨",
    "놀리다": "戏弄；让……玩；闲置；活动；挥动",
    "셈": "计算；结算；生计；情况；打算；算是",
    "묻히다": "被埋；被掩盖；陷入；埋头；沾上",
    "쌓이다": "被堆积；被砌起；积累；建立",
    "살리다": "救活；修复；发挥；唤起；鼓舞",
    "찍히다": "被戳；被盖章；被印刷；被拍摄；被认定",
    "늘리다": "增大；增加；提高；改善；延长",
    "맡다": "负责；担任；保管；占；照看；取得；闻；察觉",
    "치우다": "拿走；收拾；清理；中途停止；嫁女儿；吃完",
    "태우다": "点燃；烧焦；晒黑；抽（烟）；使焦急；承载；让乘坐；损害",
    "바": "事情；事物；方式；酒吧；吧台",
    "반죽": "面团；和面；搅拌",
    "양": "量，数量；绵羊；良；阳极；正；食量；两者，双方",
    "연": "年；风筝；缘分；姻缘；莲；联；总共；诗节，段落",
    "올": "今年；丝，缕",
    "사위": "女婿；四周，周围",
    "단속": "管束；管制，查处；断续",
    "반짝반짝": "一闪一闪；闪闪发光；灵光频现",
    "언론": "舆论；新闻媒体；传媒",
    "욕": "欲望；辱骂；责骂；侮辱；受苦",
    "본": "本，该；样本；榜样；范例",
    "채": "保持原状；栋；辆；床；还，尚；不到",
    "건": "事情，事项；件，起，份",
    "그이": "那个人；他或她；丈夫或妻子",
    "세상에": "世上，人世间；天啊",
    "호남": "湖南地区（韩国西南部，主要指全罗道）；好男儿",
    "골치": "脑袋；棘手的事，麻烦事",
    "도움말": "帮助说明；操作提示；忠告",
}

STDICT_EVIDENCE = {
    "갑": [{"word_no": 8216, "sense": "岬角"}],
    "인": [{"word_no": 269613, "sense": "磷"}],
    "동화": [{"word_no": 90004, "sense": "铜币"}],
    "침": [
        {"word_no": 491270, "sense": "针"},
        {"word_no": 338928, "sense": "针灸用针"},
    ],
    "부처": [{"word_no": 147626, "sense": "夫妻"}],
    "소요": [{"word_no": 187485, "sense": "逍遥；散步"}],
    "운동량": [{"word_no": 463228, "sense": "运动量；物理动量"}],
    "리": [{"word_no": 418967, "sense": "理由；可能性"}],
    "코스모스": [{"word_no": 490976, "sense": "宇宙"}],
}

PUBLIC_GLOSS_RESOLUTIONS = {
    "까다": "TOPIK in Depth误释为“安装”；KoreanTopik的peel与官方词典“剥”等义项一致，采用官方释义。",
    "서": "TOPIK in Depth给出无连字符的连接语尾义；KoreanTopik与官方词典均支持“西”等独立词义，不合并语尾义。",
    "실": "TOPIK in Depth的“有意义、合理”不受独立词条支持；KoreanTopik的thread与官方词典“线”等义项一致。",
    "요": "TOPIK in Depth给出礼貌语尾义；KoreanTopik与官方词典支持“褥子、这个”等独立词义，不合并语尾义。",
    "정": "TOPIK in Depth的“宇宙飞船量词”缺少官方支持；KoreanTopik的affection与官方“情”等义项一致。",
    "차로": "TOPIK in Depth把附着助词的“用车”当作词条；KoreanTopik与官方词典支持独立名词“车道”。",
    "만원": "TOPIK in Depth把应分写的“만 원（一万韩元）”并作词条；官方独立词条为“满员、满座”。",
    "세모": "TOPIK in Depth误释为“睫毛”；KoreanTopik的triangle与官方词典“三角形”一致。",
    "수상": "TOPIK in Depth混入形容词수상하다的“可疑”义；KoreanTopik与官方词典支持“获奖、首相”等名词义。",
    "이어": "TOPIK in Depth误释为“粗俗语、俚语”；KoreanTopik的then与官方词典“接着”一致。",
    "정당": "TOPIK in Depth混入形容词정당하다的“正当”义；KoreanTopik与官方词典支持名词“政党”。",
    "코트": "TOPIK in Depth把court误作“法院”；官方词典支持“大衣、球场”，不写入法院义。",
}

INITIALS = (
    "g",
    "kk",
    "n",
    "d",
    "tt",
    "r",
    "m",
    "b",
    "pp",
    "s",
    "ss",
    "",
    "j",
    "jj",
    "ch",
    "k",
    "t",
    "p",
    "h",
)
VOWELS = (
    "a",
    "ae",
    "ya",
    "yae",
    "eo",
    "e",
    "yeo",
    "ye",
    "o",
    "wa",
    "wae",
    "oe",
    "yo",
    "u",
    "wo",
    "we",
    "wi",
    "yu",
    "eu",
    "ui",
    "i",
)
FINALS = (
    "",
    "k",
    "k",
    "k",
    "n",
    "n",
    "n",
    "t",
    "l",
    "k",
    "m",
    "l",
    "l",
    "l",
    "p",
    "l",
    "m",
    "p",
    "p",
    "t",
    "t",
    "ng",
    "t",
    "t",
    "k",
    "t",
    "p",
    "t",
)


def read_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig") as file:
        return json.load(file)


def write_json(path: Path, data: dict[str, Any]) -> None:
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def compact_whitespace(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def clean_gloss(value: str) -> str:
    text = compact_whitespace(value)
    text = re.sub(r"^\d+\s*\.\s*", "", text)
    text = text.strip(" \t\r\n；;。，,")
    text = re.sub(r"的{2,}", "的", text)
    if text in {"", "(无对应词汇)", "（无对应词汇）"}:
        return ""
    return text


def official_record_glosses(record: dict[str, Any]) -> list[str]:
    glosses: list[str] = []
    headword = clean_gloss(record.get("translated_headword_zh", ""))
    if headword:
        glosses.append(headword)
    for value in record.get("translated_definitions_zh", []):
        compact = compact_whitespace(value)
        if re.match(r"^\d+\s*\.", compact):
            gloss = clean_gloss(compact)
            if gloss:
                glosses.append(gloss)
    return glosses


def compact_for_containment(value: str) -> str:
    return re.sub(r"[\s，,、；;。.!！?？（）()·…]", "", value)


def dedupe_glosses(values: Iterable[str]) -> list[str]:
    result: list[str] = []
    for raw in values:
        value = clean_gloss(raw)
        if not value:
            continue
        normalized = compact_for_containment(value)
        if not normalized:
            continue
        skip = False
        for index, existing in enumerate(result):
            existing_normalized = compact_for_containment(existing)
            if normalized == existing_normalized or normalized in existing_normalized:
                skip = True
                break
            if existing_normalized in normalized:
                result[index] = value
                skip = True
                break
        if not skip:
            result.append(value)
    return result


def normalize_public_meaning(value: str) -> str:
    text = compact_whitespace(value)
    text = re.sub(r"[（(]\s*(\d+)\s*[）)]", r" \1) ", text)
    numbered = list(re.finditer(r"(?:^|\s)(\d+)\)\s*", text))
    if numbered:
        parts: list[str] = []
        for index, match in enumerate(numbered):
            start = match.end()
            end = numbered[index + 1].start() if index + 1 < len(numbered) else len(text)
            part = clean_gloss(text[start:end])
            if part:
                parts.append(part)
        text = "；".join(parts)
    text = text.replace(";", "；")
    text = re.sub(r"；{2,}", "；", text)
    return clean_gloss(text)


def adjective_meaning(value: str) -> str:
    parts = []
    for part in value.split("；"):
        part = part.strip()
        if part and not part.endswith("的"):
            part += "的"
        if part:
            parts.append(part)
    return "；".join(parts)


def chinese_chars(value: str) -> set[str]:
    return {char for char in value if "\u4e00" <= char <= "\u9fff"}


def chinese_overlap(left: str, right: str) -> float:
    left_chars = chinese_chars(left)
    right_chars = chinese_chars(right)
    if not left_chars or not right_chars:
        return 0.0
    return len(left_chars & right_chars) / len(left_chars | right_chars)


def canonical_pos_values(raw: str) -> list[str]:
    compact = (raw or "").replace(" ", "")
    if not compact:
        return []
    values: list[str] = []
    special = {
        "의존명사": "명사",
        "보조동사": "동사",
        "보조형용사": "형용사",
    }
    for marker, canonical in special.items():
        if marker in compact and canonical not in values:
            values.append(canonical)
        compact = compact.replace(marker, "")
    for canonical in POS_CN_MAP:
        if canonical in compact and canonical not in values:
            values.append(canonical)
    if raw in LEGACY_POS_MAP and LEGACY_POS_MAP[raw] not in values:
        values.append(LEGACY_POS_MAP[raw])
    return values


def most_supported_pos(values: list[str]) -> str:
    if not values:
        return ""
    counts = Counter(values)
    return min(counts, key=lambda pos: (-counts[pos], POS_PRIORITY[pos]))


def assignment_level(entry: dict[str, Any]) -> int | None:
    match = re.search(
        r"nikl_2017_minimum_level_(\d)", entry.get("level_assignment_basis", "")
    )
    return int(match.group(1)) if match else None


def choose_pos(
    entry: dict[str, Any], records: list[dict[str, Any]]
) -> tuple[str, str, list[str]]:
    korean = entry["korean"]
    if korean in POS_OVERRIDES:
        return POS_OVERRIDES[korean], "manual_base_form_override", []

    level = assignment_level(entry)
    curriculum_values: list[str] = []
    for match in entry.get("nikl_curriculum_matches", []):
        if level is not None and match.get("grade") != f"{level}급":
            continue
        curriculum_values.extend(canonical_pos_values(match.get("pos", "")))
    if curriculum_values:
        return (
            most_supported_pos(curriculum_values),
            "nikl_2017_standard_curriculum",
            sorted(set(curriculum_values), key=POS_PRIORITY.get),
        )

    legacy_values: list[str] = []
    for match in entry.get("nikl_matches", []):
        legacy_values.extend(canonical_pos_values(match.get("pos_code", "")))
    if legacy_values:
        return (
            most_supported_pos(legacy_values),
            "nikl_2003_learning_vocabulary",
            sorted(set(legacy_values), key=POS_PRIORITY.get),
        )

    official_values: list[str] = []
    for record in records:
        official_values.extend(canonical_pos_values(record.get("pos_ko", "")))
    if official_values:
        return (
            most_supported_pos(official_values),
            "krdict_official",
            sorted(set(official_values), key=POS_PRIORITY.get),
        )

    project_values = [
        match.get("pos", "") for match in entry.get("project_vocab_matches", [])
    ]
    project_values = [value for value in project_values if value in POS_CN_MAP]
    if project_values:
        return (
            most_supported_pos(project_values),
            "project_vocab_reviewed_metadata",
            sorted(set(project_values), key=POS_PRIORITY.get),
        )

    if korean.endswith(("스럽다", "롭다", "답다", "같다", "싶다", "없다", "있다")):
        return "형용사", "suffix_heuristic", []
    if korean.endswith(("하다", "되다", "지다", "리다", "거리다", "시키다")):
        return "동사", "suffix_heuristic", []
    return "명사", "default_noun_heuristic", []


def choose_meaning(
    entry: dict[str, Any], records: list[dict[str, Any]], pos: str
) -> tuple[str, str, list[str]]:
    korean = entry["korean"]
    official_glosses = dedupe_glosses(
        gloss
        for record in records
        for gloss in official_record_glosses(record)
    )
    project_glosses = dedupe_glosses(
        match.get("meaning", "") for match in entry.get("project_vocab_matches", [])
    )
    if korean in FORCE_MEANING_OVERRIDES:
        meaning = FORCE_MEANING_OVERRIDES[korean]
        source = (
            "manual_reviewed_krdict_and_stdict_merge"
            if korean in STDICT_EVIDENCE
            else (
                "manual_reviewed_krdict_contraction_resolution"
                if korean == "얘"
                else "manual_reviewed_krdict_consolidation"
            )
        )
    elif official_glosses:
        meaning = "；".join(official_glosses)
        source = "krdict_official_chinese_headwords"
    elif korean in MEANING_OVERRIDES:
        meaning = MEANING_OVERRIDES[korean]
        source = "manual_reviewed_source_fallback"
    elif project_glosses:
        meaning = "；".join(project_glosses)
        source = "project_vocab_reviewed_metadata"
    else:
        meaning = normalize_public_meaning(entry.get("meaning_zh_source", ""))
        source = "topikindepth_chinese_fallback"

    meaning = clean_gloss(meaning)
    if pos == "형용사":
        meaning = adjective_meaning(meaning)
    return meaning, source, official_glosses


def official_pronunciation_variants(
    records: list[dict[str, Any]], selected_pos: str
) -> list[str]:
    preferred = [
        record
        for record in records
        if selected_pos in canonical_pos_values(record.get("pos_ko", ""))
    ]
    ordered_records = preferred + [record for record in records if record not in preferred]
    variants: list[str] = []
    for record in ordered_records:
        for raw in record.get("pronunciations_hangul", []):
            for value in raw.split("/"):
                cleaned = re.sub(r"[^가-힣]", "", value)
                if cleaned and cleaned not in variants:
                    variants.append(cleaned)
    return variants


def romanize_hangul(value: str) -> str:
    syllables: list[str] = []
    for index, char in enumerate(value):
        code = ord(char)
        if not 0xAC00 <= code <= 0xD7A3:
            raise ValueError(f"Unsupported pronunciation character {char!r} in {value!r}")
        offset = code - 0xAC00
        initial_index = offset // 588
        vowel_index = (offset % 588) // 28
        final_index = offset % 28
        initial = INITIALS[initial_index]
        if initial_index == 5 and index > 0:
            previous_code = ord(value[index - 1])
            if 0xAC00 <= previous_code <= 0xD7A3:
                previous_final_index = (previous_code - 0xAC00) % 28
                if FINALS[previous_final_index] == "l":
                    initial = "l"
        syllables.append(
            initial + VOWELS[vowel_index] + FINALS[final_index]
        )
    return "-".join(syllables)


class LexiconMecab:
    def __init__(self, pos_by_word: dict[str, str]):
        self.pos_by_word = pos_by_word

    def pos(self, value: str) -> list[tuple[str, str]]:
        pos = self.pos_by_word.get(value, "명사")
        tag = {
            "명사": "NNG",
            "동사": "VV",
            "형용사": "VA",
            "부사": "MAG",
            "감탄사": "IC",
            "조사": "JKS",
            "관형사": "MM",
            "대명사": "NP",
            "수사": "NR",
        }[pos]
        if pos in {"동사", "형용사"} and value.endswith("다") and len(value) > 1:
            return [(value[:-1], tag), ("다", "EF")]
        return [(value, tag)]


def load_g2p_engine(pos_by_word: dict[str, str]) -> tuple[Any, str]:
    runtime = Path(
        os.environ.get(
            "TOPIK_G2P_RUNTIME",
            str(Path(tempfile.gettempdir()) / "topik_g2p_runtime"),
        )
    )
    try:
        from g2pk2 import G2p
    except ImportError as exc:
        if runtime.exists() and str(runtime) not in sys.path:
            sys.path.insert(0, str(runtime))
        sys.modules.pop("g2pk2", None)
        try:
            from g2pk2 import G2p
        except ImportError as runtime_exc:
            raise RuntimeError(
                "g2pk2==0.0.3 is required for words without an official pronunciation. "
                "Install it into the active Python environment or into "
                f"{runtime}."
            ) from runtime_exc

    mecab = LexiconMecab(pos_by_word)
    original_check = G2p.check_mecab
    original_get = G2p.get_mecab
    try:
        G2p.check_mecab = lambda self: None
        G2p.get_mecab = lambda self: mecab
        engine = G2p()
    finally:
        G2p.check_mecab = original_check
        G2p.get_mecab = original_get
    return engine, "g2pk2_0.0.3"


def batch_label(entry: dict[str, Any]) -> str:
    basis = entry.get("level_assignment_basis", "")
    level = assignment_level(entry)
    if level:
        return f"{level}级"
    if "_B_" in basis:
        return "中级补充（2003 B）"
    if "_C_" in basis:
        return "高级补充（2003 C）"
    raise ValueError(f"Unknown assignment basis: {basis}")


def build_base_rows(
    pools: dict[str, list[dict[str, Any]]],
    krdict_entries: dict[str, list[dict[str, Any]]],
) -> dict[str, list[dict[str, Any]]]:
    rows_by_level: dict[str, list[dict[str, Any]]] = {}
    pos_by_word: dict[str, str] = {}

    for level, entries in pools.items():
        rows: list[dict[str, Any]] = []
        for entry in entries:
            korean = entry["korean"]
            records = krdict_entries.get(korean, [])
            pos, pos_source, pos_candidates = choose_pos(entry, records)
            meaning, meaning_source, official_glosses = choose_meaning(entry, records, pos)
            row = {
                "id": entry["id"],
                "level": level,
                "korean": korean,
                "pos": pos,
                "pos_cn": POS_CN_MAP[pos],
                "meaning": meaning,
                "batch": batch_label(entry),
                "source_entry": entry,
                "krdict_records": records,
                "pos_source": pos_source,
                "pos_candidates": pos_candidates,
                "meaning_source": meaning_source,
                "official_glosses": official_glosses,
                "flags": [],
            }
            rows.append(row)
            pos_by_word[korean] = pos
        rows_by_level[level] = rows

    needs_g2p = []
    for rows in rows_by_level.values():
        for row in rows:
            variants = official_pronunciation_variants(
                row["krdict_records"], row["pos"]
            )
            row["official_pronunciation_variants"] = variants
            if not variants:
                needs_g2p.append(row)

    engine = None
    engine_name = ""
    if needs_g2p:
        engine, engine_name = load_g2p_engine(pos_by_word)

    for rows in rows_by_level.values():
        for row in rows:
            variants = row["official_pronunciation_variants"]
            if variants:
                hangul_pronunciation = variants[0]
                pronunciation_source = "krdict_official"
            else:
                hangul_pronunciation = re.sub(
                    r"[^가-힣]", "", engine(row["korean"])
                )
                pronunciation_source = engine_name
            row["hangul_pronunciation"] = hangul_pronunciation
            row["pronunciation"] = romanize_hangul(hangul_pronunciation)
            row["pronunciation_source"] = pronunciation_source
            add_audit_flags(row)

    return rows_by_level


def add_audit_flags(row: dict[str, Any]) -> None:
    entry = row["source_entry"]
    records = row["krdict_records"]
    flags = row["flags"]

    if row["meaning_source"] == "manual_reviewed_source_fallback":
        flags.append("meaning_uses_manual_reviewed_fallback")
    if row["meaning_source"] == "manual_reviewed_krdict_and_stdict_merge":
        flags.append("meaning_merges_standard_dictionary_confirmed_sense")
    if row["meaning_source"] == "manual_reviewed_krdict_consolidation":
        flags.append("meaning_condenses_long_official_gloss")
    if row["meaning_source"] == "manual_reviewed_krdict_contraction_resolution":
        flags.append("meaning_resolves_contraction_homograph")
    if row["meaning_source"] == "project_vocab_reviewed_metadata":
        flags.append("meaning_reuses_reviewed_project_metadata")
    if row["meaning_source"] == "topikindepth_chinese_fallback":
        flags.append("meaning_uses_public_source_fallback")
    if len(row["official_glosses"]) > 1:
        flags.append("official_meaning_has_multiple_senses_or_homographs")
    if len(row["meaning"]) > 45:
        flags.append("meaning_longer_than_45_characters")

    public_meaning = normalize_public_meaning(entry.get("meaning_zh_source", ""))
    if (
        row["official_glosses"]
        and chinese_overlap(row["meaning"], public_meaning) < 0.08
    ):
        flags.append("official_and_public_chinese_gloss_low_character_overlap")

    if row["pos_source"].endswith("heuristic"):
        flags.append("pos_uses_heuristic")
    official_pos = sorted(
        {
            pos
            for record in records
            for pos in canonical_pos_values(record.get("pos_ko", ""))
        },
        key=POS_PRIORITY.get,
    )
    if len(official_pos) > 1:
        flags.append("krdict_has_multiple_pos")
    if row["pos_candidates"] and len(row["pos_candidates"]) > 1:
        flags.append("level_source_has_multiple_pos")

    variants = row["official_pronunciation_variants"]
    if not variants:
        flags.append("official_pronunciation_missing_uses_g2p")
    elif len(variants) > 1:
        flags.append("official_pronunciation_has_variants_primary_selected")

    source_romanization = re.sub(
        r"[^a-z]", "", entry.get("romanization_source", "").lower()
    )
    selected_romanization = row["pronunciation"].replace("-", "")
    if source_romanization and source_romanization != selected_romanization:
        flags.append("public_romanization_differs_after_sound_rules")

    project_pronunciations = {
        match.get("pronunciation", "")
        for match in entry.get("project_vocab_matches", [])
        if match.get("pronunciation")
    }
    if project_pronunciations and row["pronunciation"] not in project_pronunciations:
        flags.append("project_pronunciation_differs_from_selected")


def normalize_meaning_for_conflict(value: str) -> str:
    text = re.sub(r"[的地得了着过]", "", value)
    return compact_for_containment(text)


def distractor_conflicts(correct: str, candidate: str) -> bool:
    left = normalize_meaning_for_conflict(correct)
    right = normalize_meaning_for_conflict(candidate)
    if not left or not right or left == right:
        return True
    if min(len(left), len(right)) >= 2 and (left in right or right in left):
        return True
    left_chars = chinese_chars(left)
    right_chars = chinese_chars(right)
    if left_chars and right_chars:
        overlap = len(left_chars & right_chars) / min(len(left_chars), len(right_chars))
        if overlap >= 0.65:
            return True
    return False


def stable_candidate_order(correct_word: str, candidates: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    def key(candidate: dict[str, Any]) -> str:
        material = (
            f"topik-ii-distractor-v1|{correct_word}|{candidate['level']}|"
            f"{candidate['korean']}|{candidate['meaning']}"
        )
        return hashlib.sha256(material.encode("utf-8")).hexdigest()

    return sorted(candidates, key=key)


def load_project_words() -> list[dict[str, Any]]:
    words: list[dict[str, Any]] = []
    for level in ("beginner", "intermediate", "advanced", "topik_beginner"):
        path = DATA_DIR / f"{level}.json"
        if not path.exists():
            continue
        for word in read_json(path).get("words", []):
            if word.get("pos") not in POS_CN_MAP:
                continue
            words.append(
                {
                    "level": f"project_{level}",
                    "korean": word["korean"],
                    "pos": word["pos"],
                    "meaning": word["meaning"],
                }
            )
    return words


def assign_distractors(rows_by_level: dict[str, list[dict[str, Any]]]) -> None:
    all_topik_rows = [row for rows in rows_by_level.values() for row in rows]
    project_words = load_project_words()

    for level, rows in rows_by_level.items():
        for row in rows:
            tiers = [
                (
                    "same_book_same_pos",
                    [
                        candidate
                        for candidate in rows
                        if candidate["pos"] == row["pos"]
                    ],
                ),
                (
                    "other_topik_book_same_pos",
                    [
                        candidate
                        for candidate in all_topik_rows
                        if candidate["level"] != level and candidate["pos"] == row["pos"]
                    ],
                ),
                (
                    "project_vocab_same_pos",
                    [
                        candidate
                        for candidate in project_words
                        if candidate["pos"] == row["pos"]
                    ],
                ),
                ("same_book_any_pos", rows),
                ("other_topik_book_any_pos", all_topik_rows),
            ]

            selected: list[dict[str, str]] = []
            seen_korean = {row["korean"]}
            seen_meaning = {normalize_meaning_for_conflict(row["meaning"])}
            for tier_name, candidates in tiers:
                for candidate in stable_candidate_order(row["korean"], candidates):
                    korean = candidate["korean"]
                    meaning = candidate["meaning"]
                    normalized = normalize_meaning_for_conflict(meaning)
                    if korean in seen_korean or normalized in seen_meaning:
                        continue
                    if distractor_conflicts(row["meaning"], meaning):
                        continue
                    selected.append(
                        {
                            "meaning": meaning,
                            "korean": korean,
                            "tier": tier_name,
                        }
                    )
                    seen_korean.add(korean)
                    seen_meaning.add(normalized)
                    if len(selected) == 3:
                        break
                if len(selected) == 3:
                    break
            if len(selected) != 3:
                raise ValueError(f"Cannot generate 3 distractors for {row['korean']}")
            row["distractors"] = selected


def validate_rows(rows_by_level: dict[str, list[dict[str, Any]]]) -> list[str]:
    errors: list[str] = []
    all_korean: set[str] = set()
    beginner_words = {
        word["korean"]
        for word in read_json(DATA_DIR / "topik_beginner.json").get("words", [])
    }

    for level, rows in rows_by_level.items():
        expected_total = LEVEL_CONFIGS[level]["expected_total"]
        if len(rows) != expected_total:
            errors.append(f"{level}: expected {expected_total}, got {len(rows)}")
        expected_ids = list(range(1, len(rows) + 1))
        if [row["id"] for row in rows] != expected_ids:
            errors.append(f"{level}: IDs are not continuous")

        level_words: set[str] = set()
        for row in rows:
            required = (
                "korean",
                "pronunciation",
                "pos",
                "pos_cn",
                "meaning",
                "batch",
            )
            for key in required:
                if not row.get(key):
                    errors.append(f"{level} #{row['id']} {row['korean']}: empty {key}")
            if row["pos"] not in POS_CN_MAP:
                errors.append(f"{level} #{row['id']}: invalid POS {row['pos']}")
            if row["pos_cn"] != POS_CN_MAP.get(row["pos"]):
                errors.append(f"{level} #{row['id']}: POS translation mismatch")
            if row["korean"] in level_words:
                errors.append(f"{level}: duplicate Korean {row['korean']}")
            level_words.add(row["korean"])
            if row["korean"] in beginner_words:
                errors.append(f"{level}: overlaps TOPIK beginner {row['korean']}")
            if row["korean"] in all_korean:
                errors.append(f"cross-book duplicate Korean {row['korean']}")
            all_korean.add(row["korean"])
            if re.search(r"(?:^|\s)\d+[.)]", row["meaning"]):
                errors.append(f"{level} #{row['id']}: numbered meaning residue")
            if "的的" in row["meaning"]:
                errors.append(f"{level} #{row['id']}: duplicated adjective particle")
            if len(row["meaning"]) > 45:
                errors.append(f"{level} #{row['id']}: meaning exceeds 45 characters")
            if not re.fullmatch(r"[a-z]+(?:-[a-z]+)*", row["pronunciation"]):
                errors.append(f"{level} #{row['id']}: invalid romanization")
            if len(row["distractors"]) != 3:
                errors.append(f"{level} #{row['id']}: distractor count")
                continue
            distractor_meanings = [item["meaning"] for item in row["distractors"]]
            distractor_words = [item["korean"] for item in row["distractors"]]
            if len(set(distractor_meanings)) != 3:
                errors.append(f"{level} #{row['id']}: duplicate distractor meaning")
            if len(set(distractor_words)) != 3:
                errors.append(f"{level} #{row['id']}: duplicate distractor Korean")
            if row["meaning"] in distractor_meanings:
                errors.append(f"{level} #{row['id']}: correct meaning in distractors")
            for distractor in row["distractors"]:
                if not distractor["meaning"] or not distractor["korean"]:
                    errors.append(f"{level} #{row['id']}: empty distractor field")
    return errors


def count_values(rows: list[dict[str, Any]], key: str) -> dict[str, int]:
    return dict(sorted(Counter(row[key] for row in rows).items()))


def build_audit(
    rows_by_level: dict[str, list[dict[str, Any]]],
    krdict: dict[str, Any],
    validation_errors: list[str],
) -> dict[str, Any]:
    summary: dict[str, Any] = {}
    audit_words: dict[str, list[dict[str, Any]]] = {}
    for level, rows in rows_by_level.items():
        flag_counts = Counter(flag for row in rows for flag in row["flags"])
        distractor_tiers = Counter(
            distractor["tier"]
            for row in rows
            for distractor in row["distractors"]
        )
        summary[level] = {
            "total": len(rows),
            "meaning_sources": count_values(rows, "meaning_source"),
            "pos_sources": count_values(rows, "pos_source"),
            "pronunciation_sources": count_values(rows, "pronunciation_source"),
            "flag_counts": dict(sorted(flag_counts.items())),
            "distractor_tiers": dict(sorted(distractor_tiers.items())),
        }
        audit_words[level] = [
            {
                "id": row["id"],
                "korean": row["korean"],
                "assignment_basis": row["source_entry"]["level_assignment_basis"],
                "public_source_support_count": row["source_entry"][
                    "public_source_support_count"
                ],
                "public_sublevels": row["source_entry"]["topikindepth_sublevels"],
                "selected": {
                    "pronunciation_hangul": row["hangul_pronunciation"],
                    "pronunciation": row["pronunciation"],
                    "pos": row["pos"],
                    "pos_cn": row["pos_cn"],
                    "meaning": row["meaning"],
                    "batch": row["batch"],
                },
                "provenance": {
                    "pronunciation": row["pronunciation_source"],
                    "pos": row["pos_source"],
                    "meaning": row["meaning_source"],
                },
                "krdict_target_codes": [
                    record.get("target_code") for record in row["krdict_records"]
                ],
                "official_pronunciation_variants": row[
                    "official_pronunciation_variants"
                ],
                "official_chinese_glosses": row["official_glosses"],
                "stdict_evidence": STDICT_EVIDENCE.get(row["korean"], []),
                "public_chinese_gloss": normalize_public_meaning(
                    row["source_entry"].get("meaning_zh_source", "")
                ),
                "distractors": row["distractors"],
                "flags": row["flags"],
            }
            for row in rows
        ]

    return {
        "status": (
            "machine_validation_passed"
            if not validation_errors
            else "machine_validation_failed"
        ),
        "generated_at": TODAY,
        "scope_note": (
            "TOPIK does not publish a fixed official per-word syllabus. These books "
            "are audited preparation lists, not official TOPIK vocabulary lists."
        ),
        "krdict_snapshot": {
            "parser_version": krdict.get("parser_version"),
            "retrieved_at": krdict.get("retrieved_at"),
            "target_total": krdict.get("target_total"),
            "found_word_total": krdict.get("found_word_total"),
            "missing_word_total": len(krdict.get("missing_words", [])),
        },
        "standard_dictionary_source": {
            "name": "韩国国立国语院《标准国语大辞典》",
            "source_page": "https://stdict.korean.go.kr/",
            "usage": "复核韩国语基础词典未收录的高阶同形异义词",
            "verified_entries": STDICT_EVIDENCE,
        },
        "public_gloss_resolutions": PUBLIC_GLOSS_RESOLUTIONS,
        "manual_meaning_fallbacks": dict(sorted(MEANING_OVERRIDES.items())),
        "manual_meaning_consolidations": dict(
            sorted(FORCE_MEANING_OVERRIDES.items())
        ),
        "validation_errors": validation_errors,
        "summary": summary,
        "words": audit_words,
    }


def markdown_table_rows(values: dict[str, int]) -> str:
    return " / ".join(f"{key}: {value}" for key, value in values.items())


def meaning_source_group_count(summary: dict[str, Any], prefix: str) -> int:
    return sum(
        count
        for source, count in summary["meaning_sources"].items()
        if source.startswith(prefix)
    )


def write_audit_markdown(audit: dict[str, Any]) -> None:
    intermediate = audit["summary"]["topik_intermediate"]
    advanced = audit["summary"]["topik_advanced"]
    low_overlap_preview = []
    fallback_rows = []
    for level, words in audit["words"].items():
        for word in words:
            if "official_and_public_chinese_gloss_low_character_overlap" in word["flags"]:
                low_overlap_preview.append((level, word))
            if word["provenance"]["meaning"] == "manual_reviewed_source_fallback":
                fallback_rows.append((level, word))

    lines = [
        "# TOPIK中高级词汇元数据审计",
        "",
        f"> 创建日期：{TODAY}",
        f"> 最后更新：{TODAY}",
        f"> 状态：{'机器校验通过' if audit['status'] == 'machine_validation_passed' else '机器校验未通过'}",
        "> 用途：记录TOPIK中高级词书释义、词性、发音和测验干扰项的生成依据与异常项",
        "",
        "---",
        "",
        "## 一、口径",
        "",
        "TOPIK官方不发布固定逐词考纲。本次交付是基于国立国语院学习等级目录、"
        "韩国语基础词典及两套公开TOPIK II词表交叉审计形成的备考词书，不能称为"
        "“TOPIK官方词表”。逐词来源与分级依据见《TOPIK中高级词汇来源审计》。",
        "",
        "## 二、结果",
        "",
        "| 词书 | 词数 | 官方释义直用 | 项目释义复用 | 人工复核/合并 | 官方发音 | g2pk2补发音 |",
        "|------|-----:|-------------:|-------------:|--------------:|---------:|------------:|",
        (
            f"| TOPIK中级 | {intermediate['total']} | "
            f"{intermediate['meaning_sources'].get('krdict_official_chinese_headwords', 0)} | "
            f"{intermediate['meaning_sources'].get('project_vocab_reviewed_metadata', 0)} | "
            f"{meaning_source_group_count(intermediate, 'manual_reviewed_')} | "
            f"{intermediate['pronunciation_sources'].get('krdict_official', 0)} | "
            f"{intermediate['pronunciation_sources'].get('g2pk2_0.0.3', 0)} |"
        ),
        (
            f"| TOPIK高级 | {advanced['total']} | "
            f"{advanced['meaning_sources'].get('krdict_official_chinese_headwords', 0)} | "
            f"{advanced['meaning_sources'].get('project_vocab_reviewed_metadata', 0)} | "
            f"{meaning_source_group_count(advanced, 'manual_reviewed_')} | "
            f"{advanced['pronunciation_sources'].get('krdict_official', 0)} | "
            f"{advanced['pronunciation_sources'].get('g2pk2_0.0.3', 0)} |"
        ),
        "",
        f"- 官方词典快照命中：{audit['krdict_snapshot']['found_word_total']}/"
        f"{audit['krdict_snapshot']['target_total']}；未命中"
        f"{audit['krdict_snapshot']['missing_word_total']}词。",
        f"- 校验错误：{len(audit['validation_errors'])}。",
        f"- 中级词性来源：{markdown_table_rows(intermediate['pos_sources'])}。",
        f"- 高级词性来源：{markdown_table_rows(advanced['pos_sources'])}。",
        "",
        "## 三、校验项",
        "",
        (
            f"1. 中级{intermediate['total']}词、高级{advanced['total']}词，"
            "ID均从1连续递增。"
        ),
        "2. TOPIK初级、中级、高级三个独立词书之间韩语词形零重复。",
        "3. 所有词条的韩语、发音、词性、中文释义和批次字段非空。",
        "4. 发音优先采用韩国国立国语院官方词典实际发音；缺失项由g2pk2 0.0.3补全，"
        "再按国语罗马字规则逐音节转写。",
        "5. 每词三个干扰项优先从同词书同词性选取，并排除正确释义、重复释义和高重合释义。",
        "6. 用户可见释义最长45个字符；超长官方多义词经人工压缩后再通过校验。",
        "",
        "## 四、人工复核回退词",
        "",
        "| 词书 | 韩语 | 最终释义 |",
        "|------|------|----------|",
    ]
    for level, word in fallback_rows:
        level_cn = LEVEL_CONFIGS[level]["level_cn"]
        lines.append(f"| {level_cn} | {word['korean']} | {word['selected']['meaning']} |")

    lines.extend(
        [
            "",
            "## 五、《标准国语大辞典》补充义项",
            "",
            "下列义项在《韩国语基础词典》学习版中缺失，已通过国立国语院"
            "《标准国语大辞典》公开词条再次确认。",
            "",
            "| 韩语 | 补充义项 | 官方词条编号 |",
            "|------|----------|--------------|",
        ]
    )
    for korean, evidence_rows in STDICT_EVIDENCE.items():
        senses = "；".join(item["sense"] for item in evidence_rows)
        word_numbers = "、".join(str(item["word_no"]) for item in evidence_rows)
        lines.append(f"| {korean} | {senses} | {word_numbers} |")

    lines.extend(
        [
            "",
            "## 六、已确认的公开来源释义问题",
            "",
            "以下问题已逐项比对另一套公开词表和国立国语院官方词典；最终词书不沿用错误释义。",
            "",
            "| 韩语 | 复核结论 |",
            "|------|----------|",
        ]
    )
    for korean, note in PUBLIC_GLOSS_RESOLUTIONS.items():
        lines.append(f"| {korean} | {note} |")

    lines.extend(
        [
            "",
            "## 七、官方与公开中文词面差异样例",
            "",
            "此标记只表示中文字面字符重合度低，不等于官方释义有误。最终释义以国立国语院"
            "词典、已校对项目元数据和逐项复核结果为准；公开站点释义用于候选佐证和发现明显冲突。",
            "",
            "| 词书 | 韩语 | 官方最终释义 | 公开来源释义 |",
            "|------|------|--------------|--------------|",
        ]
    )
    for level, word in low_overlap_preview[:60]:
        level_cn = LEVEL_CONFIGS[level]["level_cn"]
        official = word["selected"]["meaning"].replace("|", "｜")
        public = word["public_chinese_gloss"].replace("|", "｜")
        lines.append(f"| {level_cn} | {word['korean']} | {official} | {public} |")

    lines.extend(
        [
            "",
            "## 八、局限",
            "",
            "1. 国立国语院课程词表是通用韩国语学习等级，不是TOPIK固定逐词考纲。",
            "2. 公开TOPIK逐词表用于考试相关候选佐证，但并非考试官方来源。",
            "3. g2pk2可覆盖规则音变，但词汇性紧音等特殊发音仍以官方词典标音优先；"
            "本次只有官方缺失项才使用规则引擎。",
            "4. 同形异义词按项目约定合并为一个学习条目，中文释义以分号区分。",
            "",
        ]
    )
    AUDIT_MD_PATH.write_text("\n".join(lines), encoding="utf-8")


def word_payload(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": row["id"],
        "korean": row["korean"],
        "pronunciation": row["pronunciation"],
        "pos": row["pos"],
        "pos_cn": row["pos_cn"],
        "meaning": row["meaning"],
    }


def write_level_json(level: str, rows: list[dict[str, Any]]) -> None:
    config = LEVEL_CONFIGS[level]
    vocab = {
        "level": level,
        "level_cn": config["level_cn"],
        "description": config["description"],
        "total": len(rows),
        "words": [word_payload(row) for row in rows],
    }
    quiz_words = []
    for row in rows:
        quiz_words.append(
            {
                **word_payload(row),
                "distractors": [
                    {
                        "meaning": distractor["meaning"],
                        "korean": distractor["korean"],
                    }
                    for distractor in row["distractors"]
                ],
            }
        )
    quiz = {
        "level": level,
        "level_cn": config["level_cn"],
        "description": config["quiz_description"],
        "total": len(rows),
        "format_note": "distractors为3个优先同词性的干扰项，正确答案为meaning字段",
        "words": quiz_words,
    }
    write_json(DATA_DIR / f"{level}.json", vocab)
    write_json(DATA_DIR / f"{level}_quiz.json", quiz)


def set_cell_shading(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_width(cell: Any, width_cm: float) -> None:
    properties = cell._tc.get_or_add_tcPr()
    width = properties.first_child_found_in("w:tcW")
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(int(Cm(width_cm).emu / 635)))
    width.set(qn("w:type"), "dxa")


def style_run(run: Any, size: float, bold: bool = False, color: str = "222222") -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def write_level_docx(level: str, rows: list[dict[str, Any]]) -> None:
    config = LEVEL_CONFIGS[level]
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.1)
    section.bottom_margin = Cm(1.1)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(f"{config['level_cn']}词汇")
    style_run(title_run, 16, bold=True, color="173F35")

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_run = note.add_run(
        f"共{len(rows)}词｜备考词书，非TOPIK官方固定考纲｜生成日期：{TODAY}"
    )
    style_run(note_run, 8.5, color="555555")

    headers = ["序号", "韩语", "发音", "词性", "中文释义", "干扰1", "干扰2", "干扰3", "批次"]
    widths = [1.1, 2.1, 2.8, 1.8, 5.3, 3.3, 3.3, 3.3, 2.6]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False

    header_cells = table.rows[0].cells
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, header in enumerate(headers):
        cell = header_cells[index]
        set_cell_width(cell, widths[index])
        set_cell_shading(cell, "DDEBE5")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        style_run(run, 8, bold=True, color="173F35")

    for row in rows:
        values = [
            str(row["id"]),
            row["korean"],
            row["pronunciation"],
            f"{row['pos']}/{row['pos_cn']}",
            row["meaning"],
            row["distractors"][0]["meaning"],
            row["distractors"][1]["meaning"],
            row["distractors"][2]["meaning"],
            row["batch"],
        ]
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            set_cell_width(cell, widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if index in {0, 1, 2, 3, 8}
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = paragraph.add_run(value)
            style_run(run, 7.2)

    document.save(str(config["docx"]))


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--write-deliverables",
        action="store_true",
        help="Write final JSON and Word files after all validation gates pass.",
    )
    args = parser.parse_args()

    krdict = read_json(KRDICT_PATH)
    if krdict.get("status") != "official_dictionary_metadata_snapshot_complete":
        raise ValueError("KRDICT metadata snapshot is incomplete")
    if krdict.get("parser_version") != 2:
        raise ValueError("KRDICT metadata snapshot must use parser version 2")

    pools: dict[str, list[dict[str, Any]]] = {}
    for level, config in LEVEL_CONFIGS.items():
        source = read_json(config["source_pool"])
        pools[level] = source["words"]

    rows_by_level = build_base_rows(pools, krdict["entries"])
    assign_distractors(rows_by_level)
    validation_errors = validate_rows(rows_by_level)
    audit = build_audit(rows_by_level, krdict, validation_errors)
    write_json(AUDIT_JSON_PATH, audit)
    write_audit_markdown(audit)

    if validation_errors:
        preview = "\n".join(validation_errors[:20])
        raise ValueError(f"TOPIK II validation failed:\n{preview}")

    if args.write_deliverables:
        for level, rows in rows_by_level.items():
            write_level_json(level, rows)
            write_level_docx(level, rows)

    print(f"status={audit['status']}")
    print(f"topik_intermediate={len(rows_by_level['topik_intermediate'])}")
    print(f"topik_advanced={len(rows_by_level['topik_advanced'])}")
    print(
        "official_dictionary_meanings="
        + str(
            sum(
                row["meaning_source"] == "krdict_official_chinese_headwords"
                for rows in rows_by_level.values()
                for row in rows
            )
        )
    )
    print(
        "official_dictionary_pronunciations="
        + str(
            sum(
                row["pronunciation_source"] == "krdict_official"
                for rows in rows_by_level.values()
                for row in rows
            )
        )
    )
    print(
        "g2p_fallback_pronunciations="
        + str(
            sum(
                row["pronunciation_source"] == "g2pk2_0.0.3"
                for rows in rows_by_level.values()
                for row in rows
            )
        )
    )
    print(f"audit_json={AUDIT_JSON_PATH}")
    print(f"audit_md={AUDIT_MD_PATH}")
    print(f"deliverables_written={args.write_deliverables}")


if __name__ == "__main__":
    main()
